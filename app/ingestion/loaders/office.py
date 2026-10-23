# import logfire
# from unstructured.partition.auto import partition

# def parse_office(file_path: str):
#     """
#     Parses Office documents (.docx, .pptx) using the Unstructured library.
#     Unlike PDFs, these formats are structured and lightweight, so they are processed locally.
#     """
#     with logfire.span("📄 Office Document Parsing", filename=file_path):
#         try:
#             # Unstructured automatically detects if it's docx or pptx
#             elements = partition(filename=file_path)
#             full_text = "\n".join([str(el) for el in elements])

#             if not full_text.strip():
#                 logfire.warning(f"⚠️ Unstructured returned empty text for {file_path}")
#             else:
#                 logfire.info(f"✅ Successfully parsed {len(full_text)} characters")

#             return full_text
#         except Exception as e:
#             logfire.error(f"❌ Office Parse Failed: {e}")
#             raise e

import logfire
from pathlib import Path


def parse_office(file_path: str) -> str:
    """Extract text from DOCX and PPTX files locally."""
    with logfire.span("Office Document Parsing", filename=file_path):
        try:
            suffix = Path(file_path).suffix.lower()

            if suffix == ".pptx":
                from pptx import Presentation

                presentation = Presentation(file_path)
                slides_text = []

                for slide_number, slide in enumerate(presentation.slides, start=1):
                    slide_parts = []

                    for shape in slide.shapes:
                        if shape.has_text_frame and shape.text.strip():
                            slide_parts.append(shape.text.strip())

                        if shape.has_table:
                            for row in shape.table.rows:
                                cells = [cell.text.strip() for cell in row.cells]
                                slide_parts.append(" | ".join(cells))

                    if slide_parts:
                        slides_text.append(
                            f"Slide {slide_number}\n" + "\n".join(slide_parts)
                        )

                return "\n\n".join(slides_text)

            elif suffix == ".docx":
                from docx import Document

                document = Document(file_path)
                paragraphs = [
                    paragraph.text.strip()
                    for paragraph in document.paragraphs
                    if paragraph.text.strip()
                ]
                return "\n".join(paragraphs)

            else:
                raise ValueError(f"Unsupported Office file: {suffix}")

        except Exception:
            logfire.exception(f"Office parsing failed: {file_path}")
            raise
